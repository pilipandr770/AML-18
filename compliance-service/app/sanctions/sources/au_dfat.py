"""Australia DFAT Consolidated Sanctions List source (XLSX workbook).

Real-file structure: one worksheet ("Consolidated List - <date>"), header in
row 1, ONE ROW PER NAME. Rows of the same listing share a numeric Reference
prefix ("2" = Primary Name, "2a" = Original Script, "3b" = Alias ...) and
repeat the listing-level columns (DOB, POB, Citizenship, Address,
Committees). Rows are therefore grouped by the leading digits of Reference.
Passport and other ID numbers exist only inside the free-text "Additional
Information" column and are NOT parsed (only the structured IMO Number column).
"""

import io
import re
from datetime import date, datetime

import requests
from openpyxl import load_workbook

from app.sanctions.sources.base import NormalizedRecord, SanctionsSource

DOWNLOAD_URL = "https://www.dfat.gov.au/sites/default/files/Australian_Sanctions_Consolidated_List.xlsx"

_MAX = 256
_REQUIRED = {
    "reference": "reference",
    "name": "name of individual or entity",
    "type": "type",
    "name_type": "name type",
}
_OPTIONAL = {
    "dob": "date of birth",
    "pob": "place of birth",
    "citizenship": "citizenship",
    "address": "address",
    "imo": "imo number",
    "committees": "committees",
    "control_date": "control date",
    "instrument": "instrument of designation",
}
_HEADER_SCAN_ROWS = 30
_TYPE_MAP = {"individual": "person", "entity": "legal_entity", "vessel": "legal_entity"}


def _cell(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).replace("\xa0", " ").strip()
    return text or None


def _find_header(rows):
    """Returns (row index, {field: column index}) of the header row, or None."""
    wanted = {**_REQUIRED, **_OPTIONAL}
    for idx, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        labels = {}
        for i, v in enumerate(row):
            label = (_cell(v) or "").lower()
            labels.setdefault(label, i)
        if _REQUIRED["reference"] in labels and _REQUIRED["name"] in labels:
            return idx, {f: labels[label] for f, label in wanted.items() if label in labels}
    return None


def _clip(value):
    return value[:_MAX] if value else None


class AustraliaDFATSource(SanctionsSource):
    name = "AU_DFAT"
    source_url = DOWNLOAD_URL

    def fetch(self) -> bytes:
        response = requests.get(
            self.source_url,
            timeout=180,
            headers={"User-Agent": "Mozilla/5.0 (AML18 sanctions fetch)"},
        )
        response.raise_for_status()
        return response.content

    def parse(self, raw: bytes):
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        rows, header = None, None
        try:
            for sheet in workbook.worksheets:
                candidate = [list(r) for r in sheet.iter_rows(values_only=True)]
                found = _find_header(candidate)
                if found:
                    rows, header = candidate, found
                    break
        finally:
            workbook.close()
        if header is None:
            raise ValueError(
                "DFAT workbook: no worksheet has a header row with 'Reference' and "
                "'Name of Individual or Entity' columns"
            )
        header_idx, cols = header
        missing = [label for f, label in _REQUIRED.items() if f not in cols]
        if missing:
            raise ValueError(f"DFAT workbook: expected columns missing: {missing}")

        def get(row, field):
            col = cols.get(field)
            return _cell(row[col]) if col is not None and col < len(row) else None

        groups = {}
        for row in rows[header_idx + 1:]:
            ref = get(row, "reference")
            if not ref:
                continue
            match = re.match(r"\d+", ref)
            groups.setdefault(match.group(0) if match else ref, []).append(row)

        for key, group in groups.items():
            record = self._build(key, group, get)
            if record is not None:
                yield record

    def _build(self, key, group, get):
        def first(field):
            return next((v for v in (get(r, field) for r in group) if v), None)

        primary_row = next(
            (r for r in group
             if (get(r, "name_type") or "").lower() == "primary name" and get(r, "name")),
            None,
        )
        if primary_row is None:
            primary_row = next((r for r in group if get(r, "name")), None)
        if primary_row is None:
            return None
        primary = get(primary_row, "name")

        seen, aliases = {primary.lower()}, []
        for r in group:
            name = get(r, "name")
            if name and name.lower() not in seen:
                seen.add(name.lower())
                aliases.append(name[:_MAX])

        entity_type = _TYPE_MAP.get((first("type") or "").lower(), "legal_entity")

        addresses, seen_addr = [], set()
        for r in group:
            addr = get(r, "address")
            if addr and addr not in seen_addr:
                seen_addr.add(addr)
                if "," in addr:
                    city, _, country = addr.rpartition(",")
                    addresses.append({"city": city.strip() or None, "country": country.strip() or None})
                else:
                    addresses.append({"city": None, "country": addr})
        country_of_residence = addresses[0]["country"] if addresses else None

        programs = []
        for r in group:
            value = get(r, "committees")
            if value and value not in programs:
                programs.append(value)
        if not programs and first("instrument"):
            programs.append(first("instrument"))

        imo = first("imo")
        return NormalizedRecord(
            source=self.name,
            source_entity_id=key,
            entity_type=entity_type,
            primary_name=primary[:_MAX],
            aliases=aliases,
            date_of_birth=first("dob") if entity_type == "person" else None,
            place_of_birth=_clip(first("pob")),
            nationality=_clip(first("citizenship")),
            country_of_residence=_clip(country_of_residence),
            national_ids=[{"type": "IMO Number", "value": imo}] if imo else [],
            addresses=addresses,
            programs=programs,
            raw_record={
                "reference": key,
                "type": first("type"),
                "control_date": first("control_date"),
                "instrument": first("instrument"),
            },
        )
